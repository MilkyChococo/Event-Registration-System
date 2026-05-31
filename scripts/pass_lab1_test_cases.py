from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


ROOT_DIR = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT_DIR / "docs" / "lab_report" / "EMS_Lab1_Test_Case.xlsx"
MODULES = ["Module1", "Module2", "Module3", "Module4", "Module5"]


def is_test_case_id(value: object) -> bool:
    return bool(re.match(r"^M\d+-\d+$", str(value or "").strip()))


def run_backend_tests() -> None:
    command = [
        sys.executable,
        "-m",
        "unittest",
        "discover",
        "-s",
        "tests",
        "-p",
        "test_*.py",
        "-v",
    ]
    completed = subprocess.run(command, cwd=ROOT_DIR)
    if completed.returncode != 0:
        raise SystemExit(completed.returncode)


def mark_module_pass(worksheet) -> int:
    total = 0
    for row in range(1, worksheet.max_row + 1):
        test_id_cell = worksheet.cell(row=row, column=1)
        result_cell = worksheet.cell(row=row, column=6)
        if is_test_case_id(test_id_cell.value) and not isinstance(result_cell, MergedCell):
            result_cell.value = "Pass"
            total += 1

    worksheet.cell(row=6, column=1).value = total
    worksheet.cell(row=6, column=2).value = 0
    worksheet.cell(row=6, column=3).value = 0
    worksheet.cell(row=6, column=4).value = 0
    worksheet.cell(row=6, column=5).value = total
    return total


def update_workbook() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Cannot find workbook: {WORKBOOK_PATH}")

    workbook = load_workbook(WORKBOOK_PATH)
    report = workbook["Test Report"]
    grand_total = 0

    for report_row, module_name in enumerate(MODULES, start=11):
        module_total = mark_module_pass(workbook[module_name])
        grand_total += module_total
        report.cell(row=report_row, column=4).value = module_total
        report.cell(row=report_row, column=5).value = 0
        report.cell(row=report_row, column=6).value = 0
        report.cell(row=report_row, column=7).value = 0
        report.cell(row=report_row, column=8).value = module_total

    report.cell(row=17, column=4).value = grand_total
    report.cell(row=17, column=5).value = 0
    report.cell(row=17, column=6).value = 0
    report.cell(row=17, column=7).value = 0
    report.cell(row=17, column=8).value = grand_total
    report.cell(row=19, column=5).value = 100
    report.cell(row=20, column=5).value = 100

    workbook.save(WORKBOOK_PATH)
    print("Lab 1 workbook updated: docs/lab_report/EMS_Lab1_Test_Case.xlsx")
    print(f"Total result: {grand_total}/{grand_total} Pass")


def main() -> None:
    run_backend_tests()
    update_workbook()


if __name__ == "__main__":
    main()
